import sys
from docx import Document

doc = Document(sys.argv[1])
print(f"Total tables: {len(doc.tables)}")
for i, table in enumerate(doc.tables):
    print(f"Table {i} has {len(table.rows)} rows")
    for r, row in enumerate(table.rows[:5]):
        print(f"  Row {r} has {len(row.cells)} cells")
        for c, cell in enumerate(row.cells):
            text = cell.text.replace('\n', '\\n')
            print(f"    Cell {c} text: {text[:30]}")
